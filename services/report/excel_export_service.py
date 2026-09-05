from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill
)
from openpyxl.utils import get_column_letter


class ExcelExportService:

    def generate_regression_report(
        self,
        report: dict
    ):

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Summary"

        self._build_summary_sheet(
            summary_sheet,
            report
        )

        methods_sheet = workbook.create_sheet(
            "Changed Methods"
        )

        self._build_methods_sheet(
            methods_sheet,
            report
        )

        scenarios_sheet = workbook.create_sheet(
            "Affected Scenarios"
        )

        self._build_scenarios_sheet(
            scenarios_sheet,
            report
        )

        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(0)

        return output

    def _build_summary_sheet(
        self,
        sheet,
        report: dict
    ):

        sheet.append(
            [
                "CodeIntelligence",
                "Regression Analysis Report"
            ]
        )

        sheet.merge_cells(
            start_row=1,
            start_column=2,
            end_row=1,
            end_column=4
        )

        sheet["A1"].font = Font(
            bold=True,
            size=16
        )

        sheet["B1"].font = Font(
            bold=True,
            size=16
        )

        sheet.append([])

        sheet.append(
            [
                "Report Status",
                report.get(
                    "status"
                )
            ]
        )

        sheet.append(
            [
                "Generated At",
                report.get(
                    "generated_at"
                )
            ]
        )

        sheet.append([])

        summary = report.get(
            "summary",
            {}
        )

        sheet.append(
            [
                "Metric",
                "Value"
            ]
        )

        self._style_header(
            sheet,
            6
        )

        sheet.append(
            [
                "Changed Java Files",
                summary.get(
                    "changed_java_files",
                    0
                )
            ]
        )

        sheet.append(
            [
                "Changed Classes",
                summary.get(
                    "changed_classes",
                    0
                )
            ]
        )

        sheet.append(
            [
                "Changed Methods",
                summary.get(
                    "changed_methods",
                    0
                )
            ]
        )

        sheet.append(
            [
                "Affected Scenarios",
                summary.get(
                    "affected_scenarios",
                    0
                )
            ]
        )

        sheet.append([])

        sheet.append(
            [
                "Changed Classes"
            ]
        )

        self._style_header(
            sheet,
            sheet.max_row
        )

        for class_name in report.get(
            "changed_classes",
            []
        ):

            sheet.append(
                [
                    class_name
                ]
            )

        self._auto_size(
            sheet
        )

    def _build_methods_sheet(
        self,
        sheet,
        report: dict
    ):

        headers = [
            "Class",
            "Method",
            "Changed Lines"
        ]

        sheet.append(
            headers
        )

        self._style_header(
            sheet,
            1
        )

        methods = report.get(
            "changed_methods",
            []
        )

        for method in methods:

            changed_lines = method.get(
                "changed_lines",
                []
            )

            sheet.append(
                [
                    method.get(
                        "class_name"
                    ),
                    method.get(
                        "method_name"
                    ),
                    ", ".join(
                        map(
                            str,
                            changed_lines
                        )
                    )
                ]
            )

        self._auto_size(
            sheet
        )

    def _build_scenarios_sheet(
        self,
        sheet,
        report: dict
    ):

        headers = [
            "Scenario ID",
            "Scenario",
            "HTTP Method",
            "Endpoint",
            "Baseline Version",
            "Impact Status",
            "Matched Classes",
            "Changed Methods",
            "Reason"
        ]

        sheet.append(
            headers
        )

        self._style_header(
            sheet,
            1
        )

        scenarios = report.get(
            "scenarios",
            []
        )

        for scenario in scenarios:

            sheet.append(
                [
                    scenario.get(
                        "scenario_id"
                    ),

                    scenario.get(
                        "scenario_code"
                    ),

                    scenario.get(
                        "http_method"
                    ),

                    scenario.get(
                        "endpoint"
                    ),

                    scenario.get(
                        "baseline_version"
                    ),

                    scenario.get(
                        "impact_status"
                    ),

                    ", ".join(
                        scenario.get(
                            "matched_classes",
                            []
                        )
                    ),

                    ", ".join(
                        scenario.get(
                            "changed_method_names",
                            []
                        )
                    ),

                    scenario.get(
                        "reason"
                    )
                ]
            )

        self._auto_size(
            sheet
        )

    def _style_header(
        self,
        sheet,
        row_number: int
    ):

        fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        for cell in sheet[
            row_number
        ]:

            cell.font = Font(
                bold=True
            )

            cell.fill = fill

            cell.alignment = Alignment(
                vertical="center"
            )

    def _auto_size(
        self,
        sheet
    ):

        for column_cells in sheet.columns:

            max_length = 0

            column_number = (
                column_cells[0]
                .column
            )

            for cell in column_cells:

                value = cell.value

                if value is None:
                    continue

                max_length = max(
                    max_length,
                    len(
                        str(value)
                    )
                )

            width = min(
                max(
                    max_length + 3,
                    12
                ),
                60
            )

            sheet.column_dimensions[
                get_column_letter(
                    column_number
                )
            ].width = width